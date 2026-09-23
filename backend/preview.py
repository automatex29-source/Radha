"""Previews for files shown in the side panel.

Office files are turned into lightweight JSON (sheet rows, slide text, document
HTML built from escaped text) so the browser can show them without plugins.
PDF, HTML, images and video are rendered by the browser from the raw bytes.
"""
import ast
import csv
import html
import io
import json
import operator
import re

MAX_PREVIEW_ROWS = 300
MAX_PREVIEW_COLS = 40
MAX_TEXT = 200_000

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_FUNC = re.compile(r"\b(SUM|AVERAGE|MIN|MAX|COUNT)\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)", re.I)
_OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv}


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n - 1


def _arith(node):
    if isinstance(node, ast.Expression):
        return _arith(node.body)
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
        return node.value
    if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
        return _OPS[type(node.op)](_arith(node.left), _arith(node.right))
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub):
        return -_arith(node.operand)
    raise ValueError("unsupported")


def _evaluate_formulas(rows: list) -> dict:
    """Compute common formulas (SUM/AVERAGE/MIN/MAX/COUNT over a range, cell refs, + - * /)
    so previews show results. Returns {(r, c): display}; anything else is left as the formula."""
    cache, results = {}, {}

    def number(r, c, depth):
        if (r, c) in cache:
            return cache[(r, c)]
        raw = rows[r][c] if r < len(rows) and c < len(rows[r]) else ""
        val = evaluate(raw, depth + 1) if raw.startswith("=") else raw
        try:
            num = float(str(val).replace(",", "").replace("$", "")) if str(val).strip() else None
        except ValueError:
            num = None
        cache[(r, c)] = num
        return num

    def evaluate(formula, depth=0):
        if depth > 20:
            raise ValueError("too deep")
        expr = formula[1:].upper()

        def func(m):
            name, c1, r1, c2, r2 = m.groups()
            vals = [number(r, c, depth) for r in range(int(r1) - 1, int(r2)) for c in range(_col_index(c1), _col_index(c2) + 1)]
            nums = [v for v in vals if v is not None]
            fn = {"SUM": sum, "MIN": min, "MAX": max, "COUNT": len,
                  "AVERAGE": lambda xs: sum(xs) / len(xs)}[name.upper()]
            return repr(float(fn(nums))) if nums or name.upper() in ("SUM", "COUNT") else "0"

        expr = _FUNC.sub(func, expr)
        expr = _REF.sub(lambda m: repr(number(int(m.group(2)) - 1, _col_index(m.group(1)), depth) or 0.0), expr)
        return _arith(ast.parse(expr, mode="eval"))

    for r, row in enumerate(rows):
        for c, raw in enumerate(row):
            if isinstance(raw, str) and raw.startswith("="):
                try:
                    val = evaluate(raw)
                    results[(r, c)] = str(int(val)) if float(val).is_integer() else f"{val:,.4f}".rstrip("0").rstrip(".")
                except Exception:
                    pass
    return results


def _xlsx(data: bytes) -> dict:
    from openpyxl import load_workbook

    # Two passes: computed values (cached by Excel) and formulas for cells we generated ourselves.
    values = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    formulas = load_workbook(io.BytesIO(data), read_only=True)
    sheets = []
    for ws in values.worksheets:
        fws = formulas[ws.title]
        rows = []
        for vrow, frow in zip(ws.iter_rows(max_row=MAX_PREVIEW_ROWS, max_col=MAX_PREVIEW_COLS, values_only=True),
                              fws.iter_rows(max_row=MAX_PREVIEW_ROWS, max_col=MAX_PREVIEW_COLS, values_only=True)):
            rows.append([_fmt(v if v is not None else f) for v, f in zip(vrow, frow)])
        while rows and not any(rows[-1]):
            rows.pop()
        width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
        rows = [r[:width] for r in rows]
        for (r, c), value in _evaluate_formulas(rows).items():
            rows[r][c] = value
        sheets.append({"name": ws.title, "rows": rows, "truncated": (ws.max_row or 0) > MAX_PREVIEW_ROWS})
    return {"type": "table", "sheets": sheets}


def _pptx(data: bytes) -> dict:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        texts, table = [], None
        title = ""
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
                if not title:
                    title = lines[0]
                    lines = lines[1:]
                texts.extend(lines)
            elif getattr(shape, "has_table", False) and shape.has_table:
                table = [[c.text for c in row.cells] for row in shape.table.rows]
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        slides.append({"index": idx, "title": title, "lines": texts, "table": table, "notes": notes})
    return {"type": "slides", "slides": slides}


def _docx(data: bytes) -> dict:
    """DOCX -> simple, safe HTML (every text node is escaped)."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    body = doc.element.body
    parts, open_list = [], None

    def runs_html(p):
        out = []
        for item in p.iter_inner_content():
            if isinstance(item, Hyperlink):
                text = html.escape(item.text)
                url = item.url or ""
                out.append(f'<a href="{html.escape(url)}">{text}</a>' if url.startswith(("http://", "https://")) else text)
                continue
            t = html.escape(item.text)
            if not t:
                continue
            if item.bold:
                t = f"<strong>{t}</strong>"
            if item.italic:
                t = f"<em>{t}</em>"
            out.append(t)
        return "".join(out)

    from docx.table import Table
    from docx.text.hyperlink import Hyperlink
    from docx.text.paragraph import Paragraph

    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            style = (p.style.name if p.style is not None else "") or ""
            list_kind = "ul" if style.startswith("List Bullet") else "ol" if style.startswith("List Number") else None
            if list_kind != open_list:
                if open_list:
                    parts.append(f"</{open_list}>")
                if list_kind:
                    parts.append(f"<{list_kind}>")
                open_list = list_kind
            content = runs_html(p)
            if list_kind:
                parts.append(f"<li>{content}</li>")
            elif style == "Title":
                parts.append(f"<h1>{content}</h1>")
            elif style.startswith("Heading"):
                level = style.rsplit(" ", 1)[-1]
                level = min(int(level), 6) if level.isdigit() else 2
                parts.append(f"<h{level}>{content}</h{level}>")
            elif style in ("Quote", "Intense Quote"):
                parts.append(f"<blockquote>{content}</blockquote>")
            elif content.strip():
                parts.append(f"<p>{content}</p>")
        elif tag == "tbl":
            if open_list:
                parts.append(f"</{open_list}>")
                open_list = None
            table = Table(child, doc)
            rows = []
            for r, row in enumerate(table.rows):
                cell_tag = "th" if r == 0 else "td"
                rows.append("<tr>" + "".join(f"<{cell_tag}>{html.escape(c.text)}</{cell_tag}>" for c in row.cells) + "</tr>")
            parts.append("<table>" + "".join(rows) + "</table>")
    if open_list:
        parts.append(f"</{open_list}>")
    return {"type": "document", "html": "\n".join(parts)}


def _csv(text: str) -> dict:
    rows = list(csv.reader(io.StringIO(text)))[:MAX_PREVIEW_ROWS]
    return {"type": "table", "sheets": [{"name": "CSV", "rows": [r[:MAX_PREVIEW_COLS] for r in rows], "truncated": False}]}


def build_preview(content_type: str, data: bytes, name: str = "") -> dict:
    ct = (content_type or "").split(";")[0].lower()
    lower = (name or "").lower()
    if ct == XLSX or lower.endswith(".xlsx"):
        return _xlsx(data)
    if ct == PPTX or lower.endswith(".pptx"):
        return _pptx(data)
    if ct == DOCX or lower.endswith(".docx"):
        return _docx(data)
    if ct == "application/pdf":
        return {"type": "pdf"}
    if ct == "text/html":
        return {"type": "html", "html": data.decode("utf-8", "replace")[:MAX_TEXT]}
    if ct.startswith("image/"):
        return {"type": "image"}
    if ct.startswith("video/"):
        return {"type": "video"}
    text = data.decode("utf-8", "replace")[:MAX_TEXT]
    if ct == "text/csv" or lower.endswith(".csv"):
        return _csv(text)
    if ct == "application/json":
        try:
            text = json.dumps(json.loads(text), indent=2)[:MAX_TEXT]
        except ValueError:
            pass
        return {"type": "code", "language": "json", "text": text}
    if ct == "text/markdown" or lower.endswith(".md"):
        return {"type": "markdown", "text": text}
    return {"type": "code", "language": "text", "text": text}
