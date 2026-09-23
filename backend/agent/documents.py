"""Document builders: Excel, PowerPoint, Word, PDF and HTML files from structured specs.

The model describes *what* goes in the file (rows, slides, markdown); these
builders own the formatting, so output looks consistent across models.
"""
import io
import os
import re
from typing import Any, List

from markdown_it import MarkdownIt

ACCENT = "6366F1"  # RADHA indigo
MAX_ROWS = 20_000
MAX_SLIDES = 60

CONTENT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "html": "text/html",
}


def safe_filename(name: str, ext: str) -> str:
    base = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", (name or "").strip())
    base = re.sub(r"[^\w\- ]+", "", base).strip().replace(" ", "-")[:80] or "document"
    return f"{base}.{ext}"


def _md():
    return MarkdownIt("commonmark").enable("table").enable("strikethrough")


# ------------------------------------------------------------------ Excel
def _cell_value(v: Any):
    if isinstance(v, (int, float, bool)) or v is None:
        return v
    s = str(v)
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d*\.\d+", s):
        return float(s)
    return s  # strings beginning with "=" are written as formulas by openpyxl


def build_xlsx(sheets: List[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not sheets:
        raise ValueError("at least one sheet is required")
    wb = Workbook()
    wb.remove(wb.active)
    for i, spec in enumerate(sheets):
        name = re.sub(r"[\[\]\*\?/\\:]", "", str(spec.get("name") or f"Sheet{i + 1}"))[:31] or f"Sheet{i + 1}"
        ws = wb.create_sheet(name)
        columns = [str(c) for c in spec.get("columns") or []]
        rows = spec.get("rows") or []
        if len(rows) > MAX_ROWS:
            raise ValueError(f"too many rows (max {MAX_ROWS})")
        if columns:
            ws.append(columns)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=ACCENT)
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
        for row in rows:
            ws.append([_cell_value(v) for v in (row if isinstance(row, list) else [row])])

        width = max(len(columns), max((len(r) for r in rows if isinstance(r, list)), default=0))
        first_data = 2 if columns else 1
        for col in range(1, width + 1):
            letter = get_column_letter(col)
            longest = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=8)
            ws.column_dimensions[letter].width = min(max(10, longest + 2), 60)
        for col_name, fmt in (spec.get("column_formats") or {}).items():
            if col_name in columns:
                letter = get_column_letter(columns.index(col_name) + 1)
                for cell in ws[letter][first_data - 1:]:
                    cell.number_format = fmt
        if columns and rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(width)}{ws.max_row}"

        chart_spec = spec.get("chart")
        if chart_spec and columns and rows:
            kind = chart_spec.get("type", "bar")
            chart = {"line": LineChart, "pie": PieChart}.get(kind, BarChart)()
            chart.title = chart_spec.get("title")
            chart.height, chart.width = 9, 18
            cat_col = columns.index(chart_spec["category_column"]) + 1 if chart_spec.get("category_column") in columns else 1
            value_cols = [columns.index(c) + 1 for c in chart_spec.get("value_columns") or [] if c in columns] or \
                         [c for c in range(1, width + 1) if c != cat_col][:1]
            for vc in value_cols:
                chart.add_data(Reference(ws, min_col=vc, min_row=1, max_row=ws.max_row), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=cat_col, min_row=2, max_row=ws.max_row))
            ws.add_chart(chart, f"{get_column_letter(width + 2)}2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------- PowerPoint
THEMES = {
    "dark": {"bg": "0B0D14", "title": "FFFFFF", "text": "D1D5DB", "accent": ACCENT, "muted": "9CA3AF"},
    "light": {"bg": "FFFFFF", "title": "111827", "text": "374151", "accent": ACCENT, "muted": "6B7280"},
}


def _bullet_items(items) -> List[tuple]:
    out = []
    for it in items or []:
        if isinstance(it, dict):
            out.append((str(it.get("text", "")), int(it.get("level", 0) or 0)))
        else:
            s = str(it)
            level = (len(s) - len(s.lstrip(" "))) // 2
            out.append((s.strip().lstrip("-•* ").strip(), min(level, 4)))
    return out


def build_pptx(slides: List[dict], title: str = "", theme: str = "dark") -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt

    if not slides:
        raise ValueError("at least one slide is required")
    if len(slides) > MAX_SLIDES:
        raise ValueError(f"too many slides (max {MAX_SLIDES})")
    t = THEMES.get(theme, THEMES["dark"])
    rgb = lambda h: RGBColor.from_string(h)  # noqa: E731

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    if title:
        prs.core_properties.title = title

    def text_box(slide, x, y, w, h, lines, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tf = slide.shapes.add_textbox(x, y, w, h).text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for i, (txt, level) in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(size * 0.45)
            run = p.add_run()
            run.text = txt
            run.font.size = Pt(size - 2 * level)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
        return tf

    def bullets(slide, x, y, w, h, items, size=22):
        lines = [(("    " * lvl) + ("•  " if lvl == 0 else "–  ") + txt, lvl) for txt, lvl in _bullet_items(items)]
        if lines:
            text_box(slide, x, y, w, h, lines, size, t["text"])

    for spec in slides:
        s = prs.slides.add_slide(blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(t["bg"])
        layout = spec.get("layout", "bullets")
        heading = str(spec.get("title") or "")

        if layout in ("title", "section"):
            text_box(s, Inches(1), Inches(2.3), W - Inches(2), Inches(1.6), [(heading, 0)], 48 if layout == "title" else 40,
                     t["title"], bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
            bar = s.shapes.add_shape(1, (W - Inches(1.2)) // 2, Inches(4.1), Inches(1.2), Inches(0.08))
            bar.fill.solid(); bar.fill.fore_color.rgb = rgb(t["accent"]); bar.line.fill.background()
            if spec.get("subtitle"):
                text_box(s, Inches(1.5), Inches(4.4), W - Inches(3), Inches(1.2), [(str(spec["subtitle"]), 0)], 22,
                         t["muted"], align=PP_ALIGN.CENTER)
        elif layout == "quote":
            text_box(s, Inches(1.5), Inches(1.8), W - Inches(3), Inches(3), [(f"“{spec.get('quote') or heading}”", 0)], 34,
                     t["title"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if spec.get("subtitle"):
                text_box(s, Inches(1.5), Inches(5), W - Inches(3), Inches(0.8), [(f"— {spec['subtitle']}", 0)], 20,
                         t["muted"], align=PP_ALIGN.CENTER)
        else:
            text_box(s, Inches(0.7), Inches(0.5), W - Inches(1.4), Inches(1), [(heading, 0)], 34, t["title"], bold=True)
            bar = s.shapes.add_shape(1, Inches(0.7), Inches(1.45), Inches(0.9), Inches(0.06))
            bar.fill.solid(); bar.fill.fore_color.rgb = rgb(t["accent"]); bar.line.fill.background()
            top, body_h = Inches(1.8), H - Inches(2.4)
            if layout == "two_column":
                col_w = (W - Inches(1.4) - Inches(0.6)) // 2
                for j, side in enumerate(("left", "right")):
                    x = Inches(0.7) + j * (col_w + Inches(0.6))
                    y = top
                    if spec.get(f"{side}_title"):
                        text_box(s, x, top, col_w, Inches(0.6), [(str(spec[f"{side}_title"]), 0)], 24, t["accent"], bold=True)
                        y = top + Inches(0.7)
                    bullets(s, x, y, col_w, body_h, spec.get(side), 20)
            elif layout == "table" and spec.get("table"):
                tbl = spec["table"]
                cols, rows = tbl.get("columns") or [], (tbl.get("rows") or [])[:18]
                n_cols = max(len(cols), max((len(r) for r in rows), default=0)) or 1
                shape = s.shapes.add_table(len(rows) + (1 if cols else 0), n_cols, Inches(0.7), top,
                                           W - Inches(1.4), Inches(0.45) * (len(rows) + 1))
                table = shape.table
                all_rows = ([cols] if cols else []) + rows
                for r, row in enumerate(all_rows):
                    for c in range(n_cols):
                        cell = table.cell(r, c)
                        cell.text = str(row[c]) if c < len(row) and row[c] is not None else ""
                        para = cell.text_frame.paragraphs[0]
                        para.runs[0].font.size = Pt(14) if para.runs else None
                        if r == 0 and cols:
                            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(t["accent"])
                            if para.runs:
                                para.runs[0].font.bold = True
                                para.runs[0].font.color.rgb = rgb("FFFFFF")
            else:
                bullets(s, Inches(0.7), top, W - Inches(1.4), body_h, spec.get("bullets"))
        if spec.get("notes"):
            s.notes_slide.notes_text_frame.text = str(spec["notes"])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------- Word
def _add_hyperlink(paragraph, url: str, text: str):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, val in (("w:color", "4F46E5"), ("w:u", "single")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        props.append(el)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _inline_runs(paragraph, inline_token):
    from docx.shared import Pt

    bold = italic = strike = False
    link = None
    link_text = []
    for tok in inline_token.children or []:
        if tok.type == "strong_open":
            bold = True
        elif tok.type == "strong_close":
            bold = False
        elif tok.type == "em_open":
            italic = True
        elif tok.type == "em_close":
            italic = False
        elif tok.type == "s_open":
            strike = True
        elif tok.type == "s_close":
            strike = False
        elif tok.type == "link_open":
            link, link_text = tok.attrs.get("href", ""), []
        elif tok.type == "link_close":
            _add_hyperlink(paragraph, link, "".join(link_text) or link)
            link = None
        elif tok.type in ("text", "code_inline", "softbreak", "hardbreak"):
            text = "\n" if tok.type == "hardbreak" else " " if tok.type == "softbreak" else tok.content
            if link is not None:
                link_text.append(text)
                continue
            run = paragraph.add_run(text)
            run.bold, run.italic, run.font.strike = bold, italic, strike
            if tok.type == "code_inline":
                run.font.name = "Consolas"
                run.font.size = Pt(10)


def build_docx(markdown: str, title: str = "") -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    if title:
        doc.core_properties.title = title
    tokens = _md().parse(markdown or "")
    lists = []  # stack of "bullet" | "ordered"
    quote = 0
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        if tok.type == "heading_open":
            level = int(tok.tag[1])
            h = doc.add_heading(level=min(level, 4) if not (level == 1 and i == 0 and title == "") else 0)
            _inline_runs(h, tokens[i + 1])
            i += 3
            continue
        if tok.type in ("bullet_list_open", "ordered_list_open"):
            lists.append("bullet" if tok.type == "bullet_list_open" else "ordered")
        elif tok.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif tok.type == "blockquote_open":
            quote += 1
        elif tok.type == "blockquote_close":
            quote -= 1
        elif tok.type == "paragraph_open":
            if lists:
                base = "List Bullet" if lists[-1] == "bullet" else "List Number"
                depth = min(len(lists), 3)
                style = base if depth == 1 else f"{base} {depth}"
            else:
                style = "Quote" if quote else None
            p = doc.add_paragraph(style=style)
            _inline_runs(p, tokens[i + 1])
            i += 3
            continue
        elif tok.type in ("fence", "code_block"):
            p = doc.add_paragraph()
            run = p.add_run(tok.content.rstrip("\n"))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        elif tok.type == "hr":
            doc.add_paragraph("─" * 40)
        elif tok.type == "table_open":
            rows, j = [], i + 1
            while tokens[j].type != "table_close":
                if tokens[j].type == "tr_open":
                    rows.append([])
                elif tokens[j].type == "inline":
                    rows[-1].append(tokens[j])
                j += 1
            n_cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=n_cols)
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, cell_tok in enumerate(row):
                    cell_p = table.cell(r, c).paragraphs[0]
                    _inline_runs(cell_p, cell_tok)
            doc.add_paragraph()
            i = j
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------------- PDF
_FONT_DIRS = ["/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "/usr/share/fonts/TTF"]
# Wide-coverage fonts used for glyphs DejaVu lacks (CJK). Extend with PDF_FALLBACK_FONTS=path1:path2.
_FALLBACK_FONTS = [
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf",
    "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf",
    "/usr/share/fonts/truetype/unifont/unifont.ttf",
]


def _find_font():
    for d in _FONT_DIRS:
        regular = os.path.join(d, "DejaVuSans.ttf")
        if os.path.exists(regular):
            return d
    return None


class _PdfRenderer:
    """Renders markdown-it tokens with fpdf2 primitives (full control over layout)."""

    INK, ACCENT, MUTED = (17, 24, 39), (79, 70, 229), (75, 85, 99)
    HEADINGS = {1: (22, INK), 2: (16, ACCENT), 3: (13.5, INK), 4: (12, INK), 5: (11, MUTED), 6: (10.5, MUTED)}

    def __init__(self, pdf, family: str, mono: str):
        self.pdf, self.family, self.mono = pdf, family, mono
        self.base_margin = pdf.l_margin

    def _font(self, style="", size=11, mono=False):
        self.pdf.set_font(self.mono if mono else self.family, style=style if not mono else "", size=size)

    def inline(self, tok, size=11, color=INK, base_style=""):
        pdf = self.pdf
        h = size * 0.52
        bold = italic = False
        link = None
        for child in tok.children or []:
            t = child.type
            if t == "strong_open": bold = True
            elif t == "strong_close": bold = False
            elif t == "em_open": italic = True
            elif t == "em_close": italic = False
            elif t == "link_open": link = child.attrs.get("href")
            elif t == "link_close": link = None
            elif t in ("softbreak",): pdf.write(h, " ")
            elif t == "hardbreak": pdf.ln(h)
            elif t in ("text", "code_inline", "image"):
                text = child.content if t != "image" else f"[image: {child.content}]"
                style = base_style + ("B" if bold else "") + ("I" if italic else "")
                style = "".join(sorted(set(style), key="BIU".index))
                if t == "code_inline":
                    self._font(size=size * 0.88, mono=True)
                    pdf.set_text_color(55, 48, 163)
                elif link:
                    self._font("U" + style.replace("U", ""), size)
                    pdf.set_text_color(*self.ACCENT)
                else:
                    self._font(style, size)
                    pdf.set_text_color(*color)
                pdf.write(h, text, link=link or "")
        pdf.set_text_color(*self.INK)

    def render(self, tokens):
        pdf = self.pdf
        lists = []  # [kind, counter]
        quote = 0
        i = 0
        while i < len(tokens):
            tok = tokens[i]
            t = tok.type
            if t == "heading_open":
                level = int(tok.tag[1])
                size, color = self.HEADINGS[level]
                pdf.ln(4 if level <= 2 else 2.5)
                pdf.set_x(pdf.l_margin)
                self.inline(tokens[i + 1], size, color, "B")
                pdf.ln(size * 0.52 + (2.5 if level <= 2 else 1.5))
                if level == 1:
                    pdf.set_draw_color(*self.ACCENT)
                    pdf.set_line_width(0.6)
                    pdf.line(pdf.l_margin, pdf.get_y() - 1, pdf.l_margin + 18, pdf.get_y() - 1)
                    pdf.ln(3)
                i += 3
                continue
            if t in ("bullet_list_open", "ordered_list_open"):
                lists.append(["ul" if t == "bullet_list_open" else "ol", int(tok.attrs.get("start", 1) or 1) - 1])
                pdf.set_left_margin(self.base_margin + 6 * len(lists) + 2 * quote)
            elif t in ("bullet_list_close", "ordered_list_close"):
                lists.pop()
                pdf.set_left_margin(self.base_margin + 6 * len(lists) + (6 if quote else 0))
                if not lists:
                    pdf.ln(1.5)
            elif t == "list_item_open":
                kind = lists[-1]
                kind[1] += 1
                marker = f"{kind[1]}." if kind[0] == "ol" else ("•" if len(lists) % 2 else "◦")
                self._font("B" if kind[0] == "ol" else "", 11)
                pdf.set_text_color(*self.ACCENT)
                pdf.set_x(pdf.l_margin - 5.5)
                pdf.cell(5.5, 5.7, marker)
                pdf.set_text_color(*self.INK)
                self._list_item_started = True
            elif t == "blockquote_open":
                quote += 1
                pdf.set_left_margin(pdf.l_margin + 6)
                self._quote_top = pdf.get_y()
            elif t == "blockquote_close":
                quote -= 1
                pdf.set_draw_color(199, 210, 254)
                pdf.set_line_width(1)
                pdf.line(pdf.l_margin - 3.5, self._quote_top, pdf.l_margin - 3.5, pdf.get_y() - 1.5)
                pdf.set_left_margin(pdf.l_margin - 6)
            elif t == "paragraph_open":
                in_item = getattr(self, "_list_item_started", False)
                if not in_item:
                    pdf.set_x(pdf.l_margin)
                self._list_item_started = False
                self.inline(tokens[i + 1], 11, self.MUTED if quote else self.INK, "I" if quote else "")
                pdf.ln(5.7 + (0.6 if lists else 2.2))
                i += 3
                continue
            elif t in ("fence", "code_block"):
                pdf.ln(1)
                self._font(size=9, mono=True)
                pdf.set_fill_color(243, 244, 246)
                pdf.set_text_color(31, 41, 55)
                pdf.multi_cell(0, 4.6, tok.content.rstrip("\n"), fill=True, padding=(2.5, 3, 2.5, 3),
                               new_x="LMARGIN", new_y="NEXT")
                pdf.set_text_color(*self.INK)
                pdf.ln(2.5)
            elif t == "hr":
                pdf.ln(2)
                pdf.set_draw_color(209, 213, 219)
                pdf.set_line_width(0.3)
                pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
                pdf.ln(4)
            elif t == "table_open":
                rows, j = [], i + 1
                while tokens[j].type != "table_close":
                    if tokens[j].type == "tr_open":
                        rows.append([])
                    elif tokens[j].type == "inline":
                        rows[-1].append(tokens[j].content)
                    j += 1
                n = max(len(r) for r in rows)
                pdf.ln(1)
                self._font(size=10)
                pdf.set_draw_color(209, 213, 219)
                pdf.set_line_width(0.2)
                from fpdf.fonts import FontFace
                with pdf.table(headings_style=FontFace(emphasis="B", color=(255, 255, 255), fill_color=(99, 102, 241)),
                               line_height=5.6, padding=1.6, text_align="LEFT",
                               cell_fill_color=(248, 250, 252), cell_fill_mode="ROWS") as table:
                    for r in rows:
                        row = table.row()
                        for c in range(n):
                            row.cell((r[c] if c < len(r) else "").replace("**", "").replace("`", ""))
                pdf.ln(3)
                i = j
            i += 1


def build_pdf(markdown: str, title: str = "") -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, margin=18)
    font_dir = _find_font()
    mono = "Courier"
    if font_dir:
        # Unicode font so non-Latin text and symbols render.
        regular = os.path.join(font_dir, "DejaVuSans.ttf")
        pdf.add_font("DejaVu", "", regular)
        for style, fname in (("B", "DejaVuSans-Bold.ttf"), ("I", "DejaVuSans-Oblique.ttf"), ("BI", "DejaVuSans-BoldOblique.ttf")):
            path = os.path.join(font_dir, fname)
            pdf.add_font("DejaVu", style, path if os.path.exists(path) else regular)
        mono_path = os.path.join(font_dir, "DejaVuSansMono.ttf")
        if os.path.exists(mono_path):
            pdf.add_font("DejaVuMono", "", mono_path)
            mono = "DejaVuMono"
        family = "DejaVu"
        fallbacks = []
        extra = [p for p in os.environ.get("PDF_FALLBACK_FONTS", "").split(":") if p]
        for i, path in enumerate(extra + _FALLBACK_FONTS):
            if os.path.exists(path) and path.lower().endswith((".ttf", ".otf")):
                try:
                    pdf.add_font(f"Fallback{i}", "", path)
                    fallbacks.append(f"Fallback{i}")
                except Exception:
                    pass
        if fallbacks:
            pdf.set_fallback_fonts(fallbacks)
    else:
        family = "Helvetica"
        # Core fonts are Latin-1 only.
        markdown = (markdown or "").replace("•", "-").encode("latin-1", "replace").decode("latin-1")
    if title:
        pdf.set_title(title)
    pdf.set_creator("RADHA by A.utomateX")
    pdf.add_page()
    renderer = _PdfRenderer(pdf, family, mono)
    renderer.render(_md().parse(markdown or ""))
    return bytes(pdf.output())
