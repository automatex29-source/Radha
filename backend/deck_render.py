"""Deck themes, slide normalisation and the PowerPoint renderer for the Decks page.

A deck is a list of slide dicts (see normalize_slide). The browser draws the same
slides on a 1280x720 canvas, and this module draws them into a 16:9 .pptx with the
same coordinates (1 px = 9525 EMU at 96 dpi), so the download matches the preview.
"""
import io
import re
import uuid
from typing import Callable, Dict, List, Optional

LAYOUTS = ("cover", "section", "bullets", "image_right", "image_left", "cards", "stats", "steps", "quote",
           "comparison", "closing")
IMAGE_LAYOUTS = ("cover", "image_right", "image_left")
MAX_SLIDES = 30

# html_* fonts are Google Fonts the preview loads; pptx_* are safe fonts PowerPoint has everywhere.
THEMES: List[dict] = [
    {"id": "aurora", "name": "Aurora", "dark": True, "bg": "0B1020", "bg2": "1E1B4B", "surface": "161B33",
     "title": "FFFFFF", "text": "CBD5E1", "muted": "94A3B8", "accent": "818CF8", "accent2": "22D3EE",
     "html_heading": "Plus Jakarta Sans", "html_body": "Inter", "pptx_heading": "Segoe UI Semibold", "pptx_body": "Segoe UI"},
    {"id": "midnight", "name": "Midnight", "dark": True, "bg": "0A0A0F", "bg2": "0A0A0F", "surface": "17171F",
     "title": "F8FAFC", "text": "D4D4D8", "muted": "A1A1AA", "accent": "F59E0B", "accent2": "FB923C",
     "html_heading": "Space Grotesk", "html_body": "Inter", "pptx_heading": "Segoe UI Semibold", "pptx_body": "Segoe UI"},
    {"id": "ocean", "name": "Ocean", "dark": True, "bg": "042F4B", "bg2": "0C4A6E", "surface": "0B3D5C",
     "title": "FFFFFF", "text": "E0F2FE", "muted": "7DD3FC", "accent": "38BDF8", "accent2": "2DD4BF",
     "html_heading": "Outfit", "html_body": "Inter", "pptx_heading": "Segoe UI Semibold", "pptx_body": "Segoe UI"},
    {"id": "sunset", "name": "Sunset", "dark": True, "bg": "2A0E2E", "bg2": "7C2D12", "surface": "3B1537",
     "title": "FFF7ED", "text": "FED7AA", "muted": "FDBA74", "accent": "FB7185", "accent2": "FBBF24",
     "html_heading": "Outfit", "html_body": "Inter", "pptx_heading": "Segoe UI Semibold", "pptx_body": "Segoe UI"},
    {"id": "forest", "name": "Forest", "dark": True, "bg": "052E16", "bg2": "14532D", "surface": "0F3D22",
     "title": "F0FDF4", "text": "D1FAE5", "muted": "86EFAC", "accent": "4ADE80", "accent2": "A3E635",
     "html_heading": "DM Serif Display", "html_body": "Inter", "pptx_heading": "Georgia", "pptx_body": "Segoe UI"},
    {"id": "paper", "name": "Paper", "dark": False, "bg": "FAF7F2", "bg2": "F3EDE3", "surface": "FFFFFF",
     "title": "1C1917", "text": "44403C", "muted": "78716C", "accent": "C2410C", "accent2": "B45309",
     "html_heading": "Playfair Display", "html_body": "Inter", "pptx_heading": "Georgia", "pptx_body": "Calibri"},
    {"id": "snow", "name": "Snow", "dark": False, "bg": "FFFFFF", "bg2": "EEF2FF", "surface": "F8FAFC",
     "title": "0F172A", "text": "334155", "muted": "64748B", "accent": "4F46E5", "accent2": "0EA5E9",
     "html_heading": "Plus Jakarta Sans", "html_body": "Inter", "pptx_heading": "Segoe UI Semibold", "pptx_body": "Segoe UI"},
    {"id": "candy", "name": "Candy", "dark": False, "bg": "FDF2F8", "bg2": "EDE9FE", "surface": "FFFFFF",
     "title": "4A044E", "text": "581C87", "muted": "9D4EDD", "accent": "DB2777", "accent2": "7C3AED",
     "html_heading": "Outfit", "html_body": "Inter", "pptx_heading": "Segoe UI Semibold", "pptx_body": "Segoe UI"},
    {"id": "mono", "name": "Mono", "dark": False, "bg": "F5F5F4", "bg2": "F5F5F4", "surface": "FFFFFF",
     "title": "0C0A09", "text": "292524", "muted": "57534E", "accent": "0C0A09", "accent2": "57534E",
     "html_heading": "Space Grotesk", "html_body": "Inter", "pptx_heading": "Segoe UI Semibold", "pptx_body": "Segoe UI"},
]
THEME_IDS = [t["id"] for t in THEMES]
DEFAULT_THEME = "aurora"


def get_theme(theme_id: str) -> dict:
    return next((t for t in THEMES if t["id"] == theme_id), THEMES[0])


# ----------------------------------------------------------- normalisation
def _s(v, limit: int) -> str:
    s = re.sub(r"\s+", " ", str(v or "")).strip()
    return s if len(s) <= limit else s[: limit - 1].rstrip() + "…"


def _strs(v, n: int, limit: int) -> List[str]:
    items = v if isinstance(v, list) else ([v] if v else [])
    out = [_s(x.get("text") if isinstance(x, dict) else x, limit) for x in items]
    return [x.lstrip("-•* ").strip() for x in out if x][:n]


def _items(v, n: int) -> List[dict]:
    out = []
    for it in (v if isinstance(v, list) else [])[:n]:
        if isinstance(it, dict):
            out.append({"icon": _s(it.get("icon"), 8), "title": _s(it.get("title"), 60), "text": _s(it.get("text"), 170)})
        elif it:
            out.append({"icon": "", "title": _s(it, 60), "text": ""})
    return out


def _stats(v) -> List[dict]:
    out = []
    for it in (v if isinstance(v, list) else [])[:4]:
        if isinstance(it, dict) and (it.get("value") or it.get("label")):
            out.append({"value": _s(it.get("value"), 12), "label": _s(it.get("label"), 80)})
    return out


def _column(v) -> dict:
    v = v if isinstance(v, dict) else {"bullets": v}
    return {"title": _s(v.get("title"), 50), "bullets": _strs(v.get("bullets"), 5, 110)}


def _image(v) -> Optional[dict]:
    if not isinstance(v, dict) or not str(v.get("url") or "").startswith("https://"):
        return None
    return {k: str(v.get(k) or "")[:500] for k in ("url", "thumb", "credit", "link")}


def normalize_slide(raw: dict, keep_id: bool = True) -> dict:
    """Clamp a model- or user-written slide to the fields and lengths the layouts can draw."""
    raw = raw if isinstance(raw, dict) else {}
    layout = str(raw.get("layout") or "bullets").strip().lower().replace("-", "_").replace(" ", "_")
    layout = {"title": "cover", "two_column": "comparison", "timeline": "steps", "image": "image_right",
              "end": "closing", "thanks": "closing"}.get(layout, layout)
    if layout not in LAYOUTS:
        layout = "bullets"
    slide = {
        "id": str(raw.get("id") or "") if keep_id and raw.get("id") else uuid.uuid4().hex[:12],
        "layout": layout,
        "title": _s(raw.get("title"), 90),
        "subtitle": _s(raw.get("subtitle"), 160),
        "body": _s(raw.get("body"), 320),
        "bullets": _strs(raw.get("bullets"), 6, 140),
        "items": _items(raw.get("items") or raw.get("cards") or raw.get("steps"), 5 if layout == "steps" else 4),
        "stats": _stats(raw.get("stats")),
        "quote": _s(raw.get("quote"), 240),
        "author": _s(raw.get("author"), 80),
        "left": _column(raw.get("left")),
        "right": _column(raw.get("right")),
        "image_query": _s(raw.get("image_query"), 60),
        "image": _image(raw.get("image")),
        "notes": str(raw.get("notes") or "")[:2000],
    }
    return slide


# --------------------------------------------------------------- PowerPoint
def _fit(size: float, text: str, soft: int) -> float:
    """Shrink a font a little for long text so it stays inside its box."""
    n = len(text or "")
    if n <= soft:
        return size
    return max(size * 0.62, size * (soft / n) ** 0.5)


def build_deck_pptx(title: str, slides: List[dict], theme_id: str,
                    image_bytes: Optional[Callable[[str], Optional[bytes]]] = None) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Emu, Pt

    t = get_theme(theme_id)
    rgb = RGBColor.from_string
    px = lambda v: Emu(int(v * 9525))  # noqa: E731
    W, H, PAD = 1280, 720, 80

    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    prs.core_properties.title = title or "Presentation"
    prs.core_properties.author = "Krish AI"
    blank = prs.slide_layouts[6]

    def rect(s, x, y, w, h, color, radius=None, line=None):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   px(x), px(y), px(w), px(h))
        if radius:
            shape.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(0.75)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def oval(s, x, y, d, color):
        shape = s.shapes.add_shape(MSO_SHAPE.OVAL, px(x), px(y), px(d), px(d))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def text(s, x, y, w, h, value, size, color, bold=False, italic=False, align="left", anchor="top",
             heading=False, spacing=1.1):
        if not value:
            return None
        box = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        p = tf.paragraphs[0]
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.line_spacing = spacing
        run = p.add_run()
        run.text = value
        f = run.font
        f.size = Pt(size * 0.75)
        f.bold, f.italic = bold, italic
        f.name = t["pptx_heading"] if heading else t["pptx_body"]
        f.color.rgb = rgb(color)
        return box

    def bullet_list(s, x, y, w, h, items, size=26):
        if not items:
            return
        size = min(size, 26 if len(items) <= 4 else 22)
        box = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(size * 0.75 * 0.7)
            p.line_spacing = 1.15
            dot = p.add_run()
            dot.text = "●  "
            dot.font.size = Pt(size * 0.75 * 0.6)
            dot.font.color.rgb = rgb(t["accent"])
            dot.font.name = t["pptx_body"]
            r = p.add_run()
            r.text = item
            r.font.size = Pt(size * 0.75)
            r.font.name = t["pptx_body"]
            r.font.color.rgb = rgb(t["text"])

    def picture(s, img: Optional[dict], x, y, w, h):
        data = image_bytes(img["url"]) if (img and image_bytes) else None
        if not data:
            # No photo: a soft accent panel keeps the layout balanced.
            panel = rect(s, x, y, w, h, t["surface"])
            panel.fill.gradient()
            panel.fill.gradient_angle = 45
            panel.fill.gradient_stops[0].color.rgb = rgb(t["accent"])
            panel.fill.gradient_stops[1].color.rgb = rgb(t["accent2"])
            return
        from PIL import Image

        try:
            iw, ih = Image.open(io.BytesIO(data)).size
        except Exception:
            return
        pic = s.shapes.add_picture(io.BytesIO(data), px(x), px(y), px(w), px(h))
        box_ratio, img_ratio = w / h, iw / ih
        if img_ratio > box_ratio:  # too wide: crop the sides
            cut = (1 - box_ratio / img_ratio) / 2
            pic.crop_left = pic.crop_right = cut
        else:  # too tall: crop top and bottom
            cut = (1 - img_ratio / box_ratio) / 2
            pic.crop_top = pic.crop_bottom = cut

    def heading(s, value, y=64, w=W - 2 * PAD, x=PAD):
        text(s, x, y, w, 70, value, _fit(44, value, 40), t["title"], bold=True, heading=True, anchor="bottom")
        rect(s, x, y + 86, 64, 6, t["accent"], radius=3)

    def background(s):
        fill = s.background.fill
        if t["bg2"] != t["bg"]:
            fill.gradient()
            fill.gradient_angle = 135
            fill.gradient_stops[0].color.rgb = rgb(t["bg"])
            fill.gradient_stops[1].color.rgb = rgb(t["bg2"])
        else:
            fill.solid()
            fill.fore_color.rgb = rgb(t["bg"])

    total = len(slides)
    for n, sl in enumerate(slides, 1):
        sl = normalize_slide(sl)
        s = prs.slides.add_slide(blank)
        background(s)
        lay = sl["layout"]

        if lay in ("cover", "closing"):
            has_img = lay == "cover" and sl["image"] is not None and image_bytes is not None
            tw = 560 if has_img else W - 2 * PAD - 120
            tx = PAD if has_img else (W - tw) / 2
            align = "left" if has_img else "center"
            if has_img:
                picture(s, sl["image"], 680, 0, 600, H)
            text(s, tx, 150, tw, 250, sl["title"], _fit(64, sl["title"], 36), t["title"], bold=True, heading=True,
                 align=align, anchor="bottom", spacing=1.0)
            rect(s, tx if has_img else (W - 90) / 2, 420, 90, 7, t["accent"], radius=3)
            text(s, tx, 452, tw, 130, sl["subtitle"] or sl["body"], _fit(26, sl["subtitle"] or sl["body"], 110),
                 t["muted"], align=align)
        elif lay == "section":
            text(s, PAD, 170, 200, 120, f"{n:02d}", 96, t["accent"], bold=True, heading=True)
            text(s, PAD, 300, W - 2 * PAD, 160, sl["title"], _fit(56, sl["title"], 40), t["title"], bold=True,
                 heading=True, spacing=1.0)
            text(s, PAD, 470, 900, 120, sl["subtitle"] or sl["body"], 26, t["muted"])
        elif lay == "bullets":
            heading(s, sl["title"])
            y = 200
            if sl["body"]:
                text(s, PAD, y, W - 2 * PAD, 80, sl["body"], _fit(24, sl["body"], 160), t["muted"])
                y += 100
            bullet_list(s, PAD, y, W - 2 * PAD, H - y - 70, sl["bullets"])
        elif lay in ("image_right", "image_left"):
            img_x = 680 if lay == "image_right" else 0
            tx = PAD if lay == "image_right" else 680
            picture(s, sl["image"], img_x, 0, 600, H)
            heading(s, sl["title"], w=520, x=tx)
            y = 200
            if sl["body"]:
                text(s, tx, y, 520, 110, sl["body"], _fit(22, sl["body"], 150), t["muted"])
                y += 130
            bullet_list(s, tx, y, 520, H - y - 60, sl["bullets"], 22)
        elif lay == "cards":
            heading(s, sl["title"])
            items = sl["items"] or [{"icon": "", "title": b, "text": ""} for b in sl["bullets"][:4]]
            k = max(1, len(items))
            gap = 28
            cw = (W - 2 * PAD - gap * (k - 1)) / k
            per_line = max(8, (cw - 56) / 9.5)  # characters per line at 18px
            longest = max((len(it["text"]) for it in items), default=0)
            ch = min(420, max(280, 196 + -(-longest // int(per_line)) * 27 + 40))
            for i, it in enumerate(items):
                cx = PAD + i * (cw + gap)
                rect(s, cx, 220, cw, ch, t["surface"], radius=20, line=t["accent"] if not t["dark"] else None)
                if it["icon"]:
                    text(s, cx + 28, 250, 80, 64, it["icon"], 44, t["accent"])
                else:
                    oval(s, cx + 28, 256, 44, t["accent"])
                text(s, cx + 28, 336, cw - 56, 70, it["title"], _fit(24, it["title"], 26), t["title"], bold=True,
                     heading=True)
                text(s, cx + 28, 410, cw - 56, ch - 200, it["text"], _fit(18, it["text"], 110), t["text"])
        elif lay == "stats":
            heading(s, sl["title"])
            stats = sl["stats"]
            k = max(1, len(stats))
            cw = (W - 2 * PAD) / k
            for i, st in enumerate(stats):
                cx = PAD + i * cw
                text(s, cx, 250, cw - 30, 120, st["value"], _fit(80, st["value"], 6), t["accent"], bold=True,
                     heading=True)
                rect(s, cx, 385, 48, 4, t["accent2"], radius=2)
                text(s, cx, 405, cw - 40, 120, st["label"], 22, t["text"])
            if sl["body"]:
                text(s, PAD, 570, W - 2 * PAD, 80, sl["body"], 20, t["muted"])
        elif lay == "steps":
            heading(s, sl["title"])
            steps = sl["items"] or [{"icon": "", "title": b, "text": ""} for b in sl["bullets"][:5]]
            k = max(1, len(steps))
            cw = (W - 2 * PAD) / k
            rect(s, PAD + 28, 298, W - 2 * PAD - cw + 0, 4, t["accent2"])
            for i, st in enumerate(steps):
                cx = PAD + i * cw
                circle = oval(s, cx, 272, 56, t["accent"])
                tf = circle.text_frame
                tf.margin_left = tf.margin_right = 0
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = str(i + 1)
                r.font.size, r.font.bold, r.font.name = Pt(18), True, t["pptx_heading"]
                r.font.color.rgb = rgb(t["bg"] if t["dark"] else "FFFFFF")
                text(s, cx, 352, cw - 30, 70, st["title"], _fit(24, st["title"], 24), t["title"], bold=True,
                     heading=True)
                text(s, cx, 425, cw - 30, 200, st["text"], _fit(18, st["text"], 110), t["text"])
        elif lay == "quote":
            text(s, PAD, 90, 200, 160, "“", 200, t["accent"], bold=True, heading=True)
            q = sl["quote"] or sl["title"]
            text(s, 160, 200, W - 320, 300, q, _fit(40, q, 110), t["title"], italic=True, heading=True,
                 align="center", anchor="middle", spacing=1.15)
            if sl["author"]:
                text(s, 160, 540, W - 320, 50, f"— {sl['author']}", 22, t["muted"], align="center")
        elif lay == "comparison":
            heading(s, sl["title"])
            cw = (W - 2 * PAD - 32) / 2
            for i, side in enumerate(("left", "right")):
                col = sl[side]
                cx = PAD + i * (cw + 32)
                rect(s, cx, 220, cw, 420, t["surface"], radius=20, line=t["accent"] if not t["dark"] else None)
                rect(s, cx, 220, cw, 8, t["accent"] if i == 0 else t["accent2"])
                text(s, cx + 36, 256, cw - 72, 50, col["title"], 28, t["accent"] if i == 0 else t["accent2"],
                     bold=True, heading=True)
                bullet_list(s, cx + 36, 330, cw - 72, 290, col["bullets"], 20)

        if lay not in ("cover", "closing", "image_right", "image_left") and total > 1:
            text(s, W - PAD - 100, H - 48, 100, 24, f"{n} / {total}", 14, t["muted"], align="right")
        notes = sl["notes"]
        if sl["image"] and sl["image"].get("credit"):
            notes = (notes + "\n\n" if notes else "") + f"Image: {sl['image']['credit']} {sl['image'].get('link', '')}"
        if notes:
            s.notes_slide.notes_text_frame.text = notes

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def public_themes() -> List[Dict]:
    return THEMES
