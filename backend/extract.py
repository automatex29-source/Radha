"""Real text extraction from uploaded documents.

No simulation: each format is parsed with a real library. Unsupported/binary
formats (e.g. images) return empty text and produce zero chunks.
"""
import csv
import io
import json

SUPPORTED_TEXT = {"txt", "md", "markdown", "json", "csv", "log"}


def extract_text(data: bytes, ext: str, content_type: str = "") -> str:
    ext = (ext or "").lower()
    try:
        if ext == "pdf":
            return _pdf(data)
        if ext in ("docx",):
            return _docx(data)
        if ext in ("xlsx", "xlsm"):
            return _xlsx(data)
        if ext == "csv":
            return _csv(data)
        if ext == "json":
            return _json(data)
        if ext in SUPPORTED_TEXT:
            return data.decode("utf-8", errors="replace")
    except Exception as exc:  # extraction failures propagate as file error status
        raise RuntimeError(f"Failed to extract {ext}: {exc}")
    return ""  # images and other binaries: no text (vision handled later)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out)


def _csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    return "\n".join(" | ".join(r) for r in rows)


def _json(data: bytes) -> str:
    obj = json.loads(data.decode("utf-8", errors="replace"))
    return json.dumps(obj, indent=2, ensure_ascii=False)


def chunk_text(text: str, size: int = 220, overlap: int = 40) -> list[str]:
    words = text.split()
    if not words:
        return []
    step = max(1, size - overlap)
    chunks = []
    for i in range(0, len(words), step):
        chunk = " ".join(words[i:i + size]).strip()
        if chunk:
            chunks.append(chunk)
    return chunks
