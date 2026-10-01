"""Offline tests for Decks: JSON parsing, slide clean-up, edit operations, generation and the .pptx renderer."""
import asyncio
import io
import json
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ.setdefault("AUTH_SECRET", "test-secret-for-unit-tests-0123456789")
os.environ["DECK_IMAGES"] = "0"

import deck_render  # noqa: E402
import decks  # noqa: E402

ALL_LAYOUTS = [
    {"layout": "cover", "title": "Solar Power", "subtitle": "Cheaper than ever"},
    {"layout": "section", "title": "Why now"},
    {"layout": "bullets", "title": "Key facts", "body": "Intro line", "bullets": ["one", "two", "three"]},
    {"layout": "image_right", "title": "Rooftops", "bullets": ["a", "b"], "image_query": "solar roof"},
    {"layout": "image_left", "title": "Farms", "body": "Big fields"},
    {"layout": "cards", "title": "Benefits", "items": [{"icon": "💡", "title": "Cheap", "text": "Low cost"},
                                                     {"icon": "🌍", "title": "Clean", "text": "No smoke"},
                                                     {"icon": "⚡", "title": "Fast", "text": "Quick build"}]},
    {"layout": "stats", "title": "Numbers", "stats": [{"value": "90%", "label": "cost drop"}, {"value": "1 TW", "label": "installed"}]},
    {"layout": "chart", "title": "Growth", "body": "Installs doubled.", "chart": {"type": "bar", "labels": ["2021", "2022", "2023"],
     "series": [{"name": "GW", "values": ["10", "14.5 GW", 21]}]}},
    {"layout": "chart", "title": "Mix", "chart": {"type": "doughnut", "labels": ["Solar", "Wind", "Coal"],
     "series": [{"name": "Share", "values": [30, 20, 50]}]}},
    {"layout": "steps", "title": "How", "items": [{"title": "Plan", "text": "x"}, {"title": "Build", "text": "y"},
                                                 {"title": "Run", "text": "z"}]},
    {"layout": "quote", "quote": "The sun is the future", "author": "Someone"},
    {"layout": "comparison", "title": "Solar vs coal", "left": {"title": "Solar", "bullets": ["clean"]},
     "right": {"title": "Coal", "bullets": ["dirty"]}},
    {"layout": "closing", "title": "Thank you", "subtitle": "Questions?"},
]


def test_parse_json_handles_fences_and_chatter():
    assert decks.parse_json('Sure!\n```json\n{"a": "x}y", "b": [1]}\n```\nthanks') == {"a": "x}y", "b": [1]}
    with pytest.raises(ValueError):
        decks.parse_json("no json here")


def test_normalize_slide_maps_aliases_and_clamps():
    s = deck_render.normalize_slide({"layout": "two-column", "title": "x" * 300, "bullets": "- only one",
                                     "image": {"url": "javascript:alert(1)"}})
    assert s["layout"] == "comparison"
    assert len(s["title"]) == 90
    assert s["bullets"] == ["only one"]
    assert s["image"] is None
    assert deck_render.normalize_slide({"layout": "weird"})["layout"] == "bullets"


@pytest.mark.parametrize("theme", deck_render.THEME_IDS)
def test_pptx_renders_every_layout_in_every_theme(theme):
    from pptx import Presentation

    data = deck_render.build_deck_pptx("Solar", ALL_LAYOUTS, theme)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == len(ALL_LAYOUTS)
    texts = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    for word in ("Solar Power", "Benefits", "90%", "Solar vs coal", "The sun is the future"):
        assert word in texts


def test_pptx_embeds_and_crops_images():
    from PIL import Image
    from pptx import Presentation

    buf = io.BytesIO()
    Image.new("RGB", (1600, 400), "red").save(buf, "JPEG")
    img = {"url": "https://example.com/a.jpg", "credit": "Photo by A"}
    slides = [{**ALL_LAYOUTS[0], "image": img}, {**ALL_LAYOUTS[3], "image": img}]
    prs = Presentation(io.BytesIO(deck_render.build_deck_pptx("T", slides, "snow", lambda url: buf.getvalue())))
    pics = [sh for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13]
    assert len(pics) == 2
    assert pics[0].crop_left > 0.2  # a very wide photo is trimmed at the sides
    assert "Photo by A" in prs.slides[0].notes_slide.notes_text_frame.text


def _deck():
    slides = [deck_render.normalize_slide(s) for s in ALL_LAYOUTS[:4]]
    return {"id": "d", "title": "T", "theme": "aurora", "slides": slides}


def test_apply_ops_uses_original_indexes():
    deck = _deck()
    ids = [s["id"] for s in deck["slides"]]
    ops = [{"op": "delete", "index": 0},
           {"op": "update", "index": 2, "slide": {"layout": "bullets", "title": "New", "bullets": ["x"]}},
           {"op": "insert", "after": 1, "slide": {"layout": "quote", "quote": "Hi"}},
           {"op": "move", "index": 3, "after": -1},
           {"op": "theme", "theme": "paper"}, {"op": "theme", "theme": "nope"},
           {"op": "title", "title": "Renamed"}, "junk", {"op": "update", "index": 99, "slide": {}}]
    slides, changed = decks.apply_ops(deck, ops)
    assert [s["id"] for s in slides][0] == ids[3]
    assert [s["layout"] for s in slides] == ["image_right", "section", "quote", "bullets"]
    assert slides[3]["id"] == ids[2] and slides[3]["title"] == "New"
    assert deck["theme"] == "paper" and deck["title"] == "Renamed"
    assert len(changed) == 1  # only the inserted slide needs a new picture


def test_compact_deck_drops_empty_fields():
    compact = json.loads(decks.compact_deck(_deck()))
    assert "id" not in compact["slides"][0] and "left" not in compact["slides"][0]
    assert compact["slides"][2]["bullets"] == ["one", "two", "three"]


def test_outline_generation_edit_and_undo(monkeypatch):
    mongomock_motor = pytest.importorskip("mongomock_motor")
    db = mongomock_motor.AsyncMongoMockClient()["t"]
    decks.init(db, "test-model")
    calls = []

    async def fake_complete(model, system, user):
        calls.append(system)
        if system == decks.OUTLINE_SYSTEM:
            return json.dumps({"title": "Solar", "slides": [{"title": f"S{i}", "points": ["p"]} for i in range(6)]})
        if system == decks.SLIDES_SYSTEM:
            n = int(user.split("exactly ")[1].split(" ")[0])
            return "```json\n" + json.dumps({"slides": [ALL_LAYOUTS[i % len(ALL_LAYOUTS)] for i in range(n)]}) + "\n```"
        return json.dumps({"reply": "Made it green.", "ops": [{"op": "theme", "theme": "forest"}]})

    monkeypatch.setattr(decks, "_complete", fake_complete)

    async def scenario():
        out = await decks.make_outline(decks.OutlineIn(prompt="solar energy", slides=6), user_id="u")
        assert out["title"] == "Solar" and len(out["outline"]) == 6
        deck = await decks.create_deck(decks.DeckIn(prompt="solar energy", outline=out["outline"], theme="snow"), user_id="u")
        assert deck["status"] == "generating"
        await asyncio.gather(*decks._tasks)
        ready = await decks.get_deck(deck["id"], user_id="u")
        assert ready["status"] == "ready" and ready["slideCount"] == 6
        assert calls.count(decks.SLIDES_SYSTEM) == 2  # 6 slides in batches of 4

        edited = await decks.edit_deck(deck["id"], decks.EditIn(instruction="make it green"), user_id="u")
        assert edited["reply"] == "Made it green." and edited["deck"]["theme"] == "forest" and edited["deck"]["canUndo"]
        undone = await decks.undo_deck(deck["id"], user_id="u")
        assert undone["theme"] == "snow"

        with pytest.raises(Exception):
            await decks.get_deck(deck["id"], user_id="someone-else")

    asyncio.run(scenario())


def test_chart_is_cleaned_and_rendered_as_a_native_chart():
    from pptx import Presentation

    ch = deck_render.normalize_slide({"layout": "graph", "chart": {"type": "column", "labels": ["a", "b"],
                                                                     "series": [{"values": ["$1,200", "x"]}]}})
    assert ch["layout"] == "chart"
    assert ch["chart"] == {"type": "bar", "labels": ["a", "b"], "series": [{"name": "Series 1", "values": [1200.0, 0.0]}],
                           "unit": ""}
    assert deck_render.normalize_slide({"layout": "chart", "chart": {"labels": ["a"], "series": [{"values": ["x"]}]}})["chart"] is None
    prs = Presentation(io.BytesIO(deck_render.build_deck_pptx("T", ALL_LAYOUTS, "ocean")))
    charts = [sh for sl in prs.slides for sh in sl.shapes if sh.has_chart]
    assert len(charts) == 2


def test_media_urls_are_allowed_as_pictures_but_other_paths_are_not():
    assert deck_render.normalize_slide({"image": {"url": "/api/media/0123456789abcdef"}})["image"]["url"].startswith("/api/")
    assert deck_render.normalize_slide({"image": {"url": "/api/decks/x/export"}})["image"] is None


def test_youtube_links_are_recognised():
    for url in ("https://www.youtube.com/watch?v=dQw4w9WgXcQ", "https://youtu.be/dQw4w9WgXcQ?t=3",
                "https://youtube.com/shorts/dQw4w9WgXcQ", "https://m.youtube.com/watch?feature=x&v=dQw4w9WgXcQ"):
        assert decks.YOUTUBE_ID.search(url).group(1) == "dQw4w9WgXcQ"
    assert decks.YOUTUBE_ID.search("https://example.com/watch?v=dQw4w9WgXcQ") is None


def test_youtube_watch_page_and_captions_are_parsed():
    import youtube_source

    html = ('<meta name="title" content="How solar works">'
            '"shortDescription":"Panels turn light into power.\\nMore below","captionTracks":['
            '{"baseUrl":"https://www.youtube.com/api/timedtext?v=x&lang=fr","languageCode":"fr"},'
            '{"baseUrl":"https://www.youtube.com/api/timedtext?v=x&lang=en&kind=asr","languageCode":"en","kind":"asr"},'
            '{"baseUrl":"https://www.youtube.com/api/timedtext?v=x&lang=en","languageCode":"en"}],"x":1')
    page = youtube_source.parse_watch_page(html)
    assert page["title"] == "How solar works"
    assert page["description"] == "Panels turn light into power.\nMore below"
    assert youtube_source.pick_track(page["tracks"]).endswith("lang=en")
    assert youtube_source.parse_json3({"events": [{"segs": [{"utf8": "Hello"}, {"utf8": " there"}]}, {}]}) == "Hello there"
    assert youtube_source.parse_xml_captions('<text start="0">Tom &amp; Jerry</text><text>run</text>') == "Tom & Jerry run"


def test_video_without_captions_still_makes_a_source(monkeypatch):
    import youtube_source

    async def title_only(video_id):
        return {"title": "Future of electric cars", "channel": "Tech Talks", "description": "", "transcript": ""}

    monkeypatch.setattr(youtube_source, "read", title_only)
    got = asyncio.run(decks.read_video("https://youtu.be/dQw4w9WgXcQ", "dQw4w9WgXcQ"))
    assert "Future of electric cars" in got["text"] and "Show transcript" in got["note"]

    async def nothing(video_id):
        return {"title": "", "channel": "", "description": "", "transcript": ""}

    monkeypatch.setattr(youtube_source, "read", nothing)
    with pytest.raises(decks.HTTPException) as err:
        asyncio.run(decks.read_video("https://youtu.be/dQw4w9WgXcQ", "dQw4w9WgXcQ"))
    assert "Paste text" in err.value.detail


def test_stock_photos_are_ranked_by_matching_words():
    import deck_images

    good = {"_text": "farmer checking wheat field India", "_width": 2000}
    big_but_wrong = {"_text": "city skyline night", "_width": 4000}
    assert deck_images.score("farmer wheat field", good) > deck_images.score("farmer wheat field", big_but_wrong)
    assert deck_images.queries_for({"image_query": "farmer checking wheat field", "title": "Better harvests ahead"},
                                   "Indian agriculture") == ["farmer checking wheat field", "farmer checking",
                                                             "Better harvests ahead", "Indian agriculture"]


def test_pictures_fill_in_background_and_fall_back_to_photos(monkeypatch):
    mongomock_motor = pytest.importorskip("mongomock_motor")
    import deck_images

    db = mongomock_motor.AsyncMongoMockClient()["t"]
    decks.init(db, "test-model")
    monkeypatch.setattr(deck_images, "enabled", lambda: True)
    made = []

    async def fake_ai(db_, user_id, slide, topic="", prompt=None):
        made.append(slide["id"])
        if len(made) == 1:
            raise RuntimeError("busy")
        return {"url": "/api/media/aaaaaaaaaaaa", "thumb": "", "credit": "AI-made picture", "link": ""}

    async def fake_stock(slide, used, topic=""):
        return {"url": "https://img.test/photo.jpg", "thumb": "", "credit": "Photo", "link": ""}

    monkeypatch.setattr(deck_images, "make_ai", fake_ai)
    monkeypatch.setattr(deck_images, "find_stock", fake_stock)
    slides = [deck_render.normalize_slide(s) for s in (ALL_LAYOUTS[0], ALL_LAYOUTS[1], ALL_LAYOUTS[3])]

    async def scenario():
        await db.decks.insert_one({"id": "d", "userId": "u", "title": "Solar", "theme": "aurora", "pictures": "ai",
                                   "slides": slides, "status": "ready", "picturesPending": True, "updatedAt": decks._now()})
        await decks.fill_pictures("d")
        doc = await db.decks.find_one({"id": "d"})
        assert [(s["image"] or {}).get("url") for s in doc["slides"]] == \
            ["https://img.test/photo.jpg", None, "/api/media/aaaaaaaaaaaa"]
        assert doc["picturesPending"] is False

    asyncio.run(scenario())


def test_import_url_reads_pages_and_rephrase_returns_brief(monkeypatch):
    from agent import web

    async def fake_page(url):
        return {"url": url, "title": "Solar report", "text": "Solar grew fast in 2023. " * 10, "links": []}

    async def fake_complete(model, system, user):
        return "Topic: Solar\nAudience: students"

    monkeypatch.setattr(web, "fetch_page", fake_page)
    monkeypatch.setattr(decks, "_complete", fake_complete)

    async def scenario():
        out = await decks.import_url(decks.UrlIn(url="example.com/report"), user_id="u")
        assert out["kind"] == "link" and out["title"] == "Solar report" and out["url"] == "https://example.com/report"
        brief = await decks.rephrase_prompt(decks.RephraseIn(prompt="solar ppt for kids"), user_id="u")
        assert brief["prompt"].startswith("Topic: Solar")

    asyncio.run(scenario())


def test_ai_pictures_stop_after_repeated_failures(monkeypatch):
    mongomock_motor = pytest.importorskip("mongomock_motor")
    import deck_images

    db = mongomock_motor.AsyncMongoMockClient()["t"]
    decks.init(db, "test-model")
    monkeypatch.setattr(deck_images, "enabled", lambda: True)
    tries = []

    async def broke_ai(*a, **k):
        tries.append(1)
        raise RuntimeError("HTTP 402: Insufficient balance")

    async def fake_stock(slide, used, topic=""):
        return {"url": f"https://img.test/{slide['id']}.jpg", "thumb": "", "credit": "Photo", "link": ""}

    monkeypatch.setattr(deck_images, "make_ai", broke_ai)
    monkeypatch.setattr(deck_images, "find_stock", fake_stock)
    slides = [deck_render.normalize_slide({"layout": "image_right", "title": f"S{i}", "image_query": "sun"}) for i in range(5)]

    async def scenario():
        await db.decks.insert_one({"id": "d", "userId": "u", "title": "T", "theme": "aurora", "pictures": "ai",
                                   "slides": slides, "status": "ready", "picturesPending": True, "updatedAt": decks._now()})
        await decks.fill_pictures("d")
        doc = await db.decks.find_one({"id": "d"})
        assert all(s["image"]["url"].startswith("https://img.test/") for s in doc["slides"])
        assert len(tries) == 2

    asyncio.run(scenario())
