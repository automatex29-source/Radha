"""Offline tests for document builders, previews, video providers, heartbeats and the browser tool."""
import asyncio
import http.server
import io
import socketserver
import sys
import threading
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import preview  # noqa: E402
from agent import Tool, ToolContext, ToolOutput, ToolRegistry, documents, run_agent, runtime, video, web  # noqa: E402

MD = """# Report
Intro with **bold**, *italic*, `code` and a [link](https://example.com). café ✓

## Section
- one
  - nested
- two

1. first
2. second

> quoted

| Region | Revenue |
|---|---|
| EMEA | 120 |

```python
print("hi")
```
"""


def run(coro):
    return asyncio.run(coro)


class TestDocuments:
    def test_xlsx_values_formulas_formats_and_chart(self):
        from openpyxl import load_workbook

        data = documents.build_xlsx([{
            "name": "Sales/Q1", "columns": ["Month", "Revenue"],
            "rows": [["Jan", "100"], ["Feb", "120.5"], ["Total", "=SUM(B2:B3)"]],
            "column_formats": {"Revenue": "$#,##0"},
            "chart": {"type": "line", "title": "Rev", "category_column": "Month", "value_columns": ["Revenue"]},
        }])
        ws = load_workbook(io.BytesIO(data))["SalesQ1"]  # invalid sheet-name chars stripped
        assert [c.value for c in ws[1]] == ["Month", "Revenue"]
        assert ws["B2"].value == 100 and ws["B3"].value == 120.5 and ws["B4"].value == "=SUM(B2:B3)"
        assert ws["B2"].number_format == "$#,##0"
        assert ws.freeze_panes == "A2" and len(ws._charts) == 1

    def test_xlsx_requires_sheets(self):
        with pytest.raises(ValueError):
            documents.build_xlsx([])

    def test_pptx_all_layouts(self):
        from pptx import Presentation

        data = documents.build_pptx([
            {"layout": "title", "title": "Deck", "subtitle": "Sub"},
            {"layout": "bullets", "title": "Points", "bullets": ["a", "  b"], "notes": "speak"},
            {"layout": "two_column", "title": "Compare", "left": ["x"], "right": ["y"]},
            {"layout": "table", "title": "T", "table": {"columns": ["A", "B"], "rows": [["1", "2"]]}},
            {"layout": "quote", "title": "Q", "quote": "Be bold", "subtitle": "Me"},
            {"layout": "section", "title": "Part 2"},
        ], title="Deck", theme="light")
        prs = Presentation(io.BytesIO(data))
        assert len(prs.slides) == 6
        texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
        assert "Points" in texts and any("a" in t and "b" in t for t in texts)
        assert prs.slides[1].notes_slide.notes_text_frame.text == "speak"

    def test_docx_structure(self):
        from docx import Document

        doc = Document(io.BytesIO(documents.build_docx(MD, "Report")))
        styles = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
        assert ("Heading 1", "Report") in styles and ("Heading 2", "Section") in styles
        assert ("List Bullet 2", "nested") in styles and ("List Number", "first") in styles
        assert ("Quote", "quoted") in styles
        assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [["Region", "Revenue"], ["EMEA", "120"]]

    def test_pdf_renders_text(self):
        from pypdf import PdfReader

        text = PdfReader(io.BytesIO(documents.build_pdf(MD, "Report"))).pages[0].extract_text()
        for expected in ("Report", "bold", "nested", "second", "quoted", "EMEA", 'print("hi")'):
            assert expected in text

    def test_safe_filename(self):
        assert documents.safe_filename("../../etc/My Report!.pdf", "pdf") == "etcMy-Report.pdf"
        assert documents.safe_filename("", "xlsx") == "document.xlsx"


class TestPreview:
    def test_xlsx_preview_trims_empty_columns_and_computes_formulas(self):
        data = documents.build_xlsx([{"name": "S", "columns": ["A", "B"],
                                      "rows": [["1", "=A2*2"], ["4", "=B2+A3"], ["=SUM(A2:A3)", "=AVERAGE(B2:B3)"],
                                               ["x", "=VLOOKUP(1,A2:B3,2)"]],
                                      "chart": {"type": "bar"}}])
        p = preview.build_preview(documents.CONTENT_TYPES["xlsx"], data)
        assert p == {"type": "table", "sheets": [{"name": "S", "rows": [
            ["A", "B"], ["1", "2"], ["4", "6"], ["5", "4"], ["x", "=VLOOKUP(1,A2:B3,2)"]], "truncated": False}]}

    def test_pptx_preview(self):
        data = documents.build_pptx([{"layout": "bullets", "title": "Hello", "bullets": ["x"], "notes": "n"}])
        slide = preview.build_preview(documents.CONTENT_TYPES["pptx"], data)["slides"][0]
        assert slide["title"] == "Hello" and slide["notes"] == "n" and any("x" in l for l in slide["lines"])

    def test_docx_preview_is_escaped_html(self):
        from docx import Document

        html = preview.build_preview(documents.CONTENT_TYPES["docx"],
                                     documents.build_docx("# T\n\n[site](https://a.b)\n\n- item", ""))["html"]
        assert '<a href="https://a.b">site</a>' in html and "<li>item</li>" in html
        # Text from an uploaded/hostile .docx must come out escaped.
        doc = Document()
        doc.add_paragraph("Hello <script>alert(1)</script>")
        buf = io.BytesIO()
        doc.save(buf)
        html = preview.build_preview(documents.CONTENT_TYPES["docx"], buf.getvalue())["html"]
        assert "<script>" not in html and "&lt;script&gt;alert(1)&lt;/script&gt;" in html

    def test_simple_types(self):
        assert preview.build_preview("application/pdf", b"%PDF")["type"] == "pdf"
        assert preview.build_preview("text/csv", b"a,b\n1,2")["sheets"][0]["rows"] == [["a", "b"], ["1", "2"]]
        assert preview.build_preview("application/json", b'{"a":1}')["text"] == '{\n  "a": 1\n}'
        assert preview.build_preview("text/html", b"<b>x</b>") == {"type": "html", "html": "<b>x</b>"}


class TestVideo:
    def test_provider_selection(self, monkeypatch):
        for k in ("VIDEO_PROVIDER", "OPENAI_API_KEY", "GEMINI_API_KEY", "GOOGLE_API_KEY"):
            monkeypatch.delenv(k, raising=False)
        assert video.provider() is None and not video.available()
        monkeypatch.setenv("GEMINI_API_KEY", "g")
        assert video.provider() == "gemini"
        monkeypatch.setenv("OPENAI_API_KEY", "o")
        assert video.provider() == "openai"
        monkeypatch.setenv("VIDEO_PROVIDER", "gemini")
        assert video.provider() == "gemini"

    def test_sora_submit_poll_download(self, monkeypatch):
        monkeypatch.setenv("OPENAI_API_KEY", "k")
        monkeypatch.delenv("VIDEO_PROVIDER", raising=False)
        monkeypatch.setattr(video, "POLL_SECONDS", 0)
        seen, polls = [], iter(["in_progress", "completed"])

        def handler(req):
            seen.append((req.method, req.url.path))
            if req.method == "POST":
                body = req.content.decode()
                assert 'name="seconds"\r\n\r\n12' in body and "720x1280" in body and "sora-2" in body
                return httpx.Response(200, json={"id": "v1", "status": "queued"})
            if req.url.path.endswith("/content"):
                return httpx.Response(200, content=b"MP4")
            return httpx.Response(200, json={"id": "v1", "status": next(polls)})

        monkeypatch.setattr(video, "_http_client", lambda: httpx.AsyncClient(
            base_url="https://api.openai.com/v1", transport=httpx.MockTransport(handler)))
        assert run(video.generate_video("cat", 11, "portrait")) == (b"MP4", "video/mp4")
        assert seen[-1] == ("GET", "/v1/videos/v1/content") and seen.count(("GET", "/v1/videos/v1")) == 2

    def test_sora_failure_surfaces_message(self, monkeypatch):
        monkeypatch.setenv("OPENAI_API_KEY", "k")
        monkeypatch.delenv("VIDEO_PROVIDER", raising=False)
        monkeypatch.setattr(video, "POLL_SECONDS", 0)

        def handler(req):
            if req.method == "POST":
                return httpx.Response(200, json={"id": "v1", "status": "queued"})
            return httpx.Response(200, json={"id": "v1", "status": "failed", "error": {"message": "moderation"}})

        monkeypatch.setattr(video, "_http_client", lambda: httpx.AsyncClient(
            base_url="https://x/v1", transport=httpx.MockTransport(handler)))
        with pytest.raises(video.VideoError, match="moderation"):
            run(video.generate_video("x"))

    def test_veo_polls_operation(self, monkeypatch):
        monkeypatch.setenv("VIDEO_PROVIDER", "gemini")
        monkeypatch.setattr(video, "POLL_SECONDS", 0)
        done_op = SimpleNamespace(done=True, error=None, response=SimpleNamespace(
            generated_videos=[SimpleNamespace(video=SimpleNamespace(video_bytes=None))]))
        calls = {}

        class FakeAio:
            class models:
                @staticmethod
                async def generate_videos(**kw):
                    calls["config"] = kw["config"]
                    return SimpleNamespace(done=False)

            class operations:
                @staticmethod
                async def get(op):
                    return done_op

            class files:
                @staticmethod
                async def download(file):
                    return b"VEO"

        monkeypatch.setattr(video, "_genai_client", lambda: SimpleNamespace(aio=FakeAio))
        assert run(video.generate_video("sunset", orientation="portrait")) == (b"VEO", "video/mp4")
        assert calls["config"].aspect_ratio == "9:16"


class TestHeartbeat:
    def test_slow_tool_emits_heartbeats(self, monkeypatch):
        monkeypatch.setattr(runtime, "HEARTBEAT_SECONDS", 0.05)

        async def slow(ctx, args):
            await asyncio.sleep(0.2)
            return ToolOutput(content="ok")

        reg = ToolRegistry().register(Tool("slow", "Slow", {"type": "object"}, slow))
        turns = iter([[{"type": "tool_calls", "calls": [{"id": "1", "name": "slow", "arguments": "{}"}]}],
                      [{"type": "text", "text": "done"}]])

        async def llm(model, messages, tools):
            for ev in next(turns):
                yield ev

        async def collect():
            return [e["type"] async for e in run_agent(llm, reg, ToolContext(None, "u"), "m", [])]

        types = run(collect())
        assert types[0] == "tool_start" and types[-2:] == ["tool_end", "text"]
        assert types.count("heartbeat") >= 2


class TestBrowser:
    @pytest.fixture
    def site(self, monkeypatch):
        hits = []

        class Page(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                if self.path == "/next":
                    body = b"<title>Next</title><h1>Arrived</h1>"
                else:
                    body = (b"<title>Home</title><input id=q placeholder='Search'><button onclick=\""
                            b"document.getElementById('o').innerText='You typed '+document.getElementById('q').value\">Go</button>"
                            b"<p id=o></p><a href='/next'>Next page</a>"
                            b"<img src='http://localhost:%d/leak'>" % internal.server_address[1])
                self.send_response(200)
                self.send_header("Content-Type", "text/html")
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, *a):
                pass

        class Internal(Page):
            def do_GET(self):
                hits.append(self.path)
                self.send_response(200)
                self.end_headers()

        internal = socketserver.TCPServer(("127.0.0.1", 0), Internal)
        public = socketserver.TCPServer(("127.0.0.1", 0), Page)
        for srv in (internal, public):
            threading.Thread(target=srv.serve_forever, daemon=True).start()
        port = public.server_address[1]
        original = web.assert_public_url

        async def allow_site(url):  # treat the test site as "public"
            if f"127.0.0.1:{port}" not in url:
                await original(url)

        monkeypatch.setattr(web, "assert_public_url", allow_site)
        yield f"http://127.0.0.1:{port}/", hits
        for srv in (internal, public):
            srv.shutdown()

    def test_open_type_click_and_guard(self, site):
        from agent import browser
        pytest.importorskip("playwright")
        url, hits = site
        mgr = browser.BrowserManager()

        async def scenario():
            try:
                snap = await mgr.act("k", "open", {"url": url})
            except browser.BrowserError as exc:
                pytest.skip(f"no browser available: {exc}")
            refs = {e["label"]: e["ref"] for e in snap["elements"]}
            assert snap["title"] == "Home" and snap["screenshot"][:2] == b"\xff\xd8"
            await mgr.act("k", "type", {"ref": refs["Search"], "text": "shoes"})
            snap = await mgr.act("k", "click", {"ref": refs["Go"]})
            assert "You typed shoes" in snap["text"]
            snap = await mgr.act("k", "click", {"ref": refs["Next page"]})
            assert snap["title"] == "Next"
            with pytest.raises(web.FetchError):
                await mgr.act("k", "open", {"url": "http://169.254.169.254/"})
            await mgr.shutdown()

        run(scenario())
        assert hits == []  # the page's request to an internal host was blocked
        assert "[1]" in browser.describe({"url": url, "title": "t", "text": "x", "elements": [
            {"ref": 1, "kind": "button", "label": "Go", "inView": True}]})
